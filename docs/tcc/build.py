"""Monta o TCC (docx e pdf) a partir dos capítulos em Markdown.

Uso: python docs/tcc/build.py [--out DIR]

Passos: numera títulos, reordena legendas de figuras, gera sumário e listas com
números de página (lidos de um PDF de primeira passada), roda o pandoc com um
modelo ABNT e ajusta tabelas e cabeçalho com python-docx. Requer pandoc, LibreOffice
(soffice) e pdftotext.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
FONTE = "Arial"
LARGURA_TWIPS = 9071  # 16 cm de texto (A4 com margens de 3 e 2 cm)

CAPITULOS = [
    "01-introducao.md",
    "02-revisao.md",
    "03-fundamentacao.md",
    "04-materiais-metodos.md",
    "05-resultados.md",
    "06-consideracoes.md",
]
POS = ["07-referencias.md", "08-apendices.md"]

SIGLAS = [
    ("ABNT", "Associação Brasileira de Normas Técnicas"),
    ("API", "Application Programming Interface (Interface de Programação de Aplicações)"),
    ("ASGI", "Asynchronous Server Gateway Interface"),
    ("BS", "Block Size (tamanho do bloco, no ISO-TP)"),
    ("CAN", "Controller Area Network (Rede de Área do Controlador)"),
    ("CF", "Consecutive Frame (quadro consecutivo, no ISO-TP)"),
    ("CRC", "Cyclic Redundancy Check (Verificação Cíclica de Redundância)"),
    ("CSV", "Comma-Separated Values"),
    ("DLC", "Data Length Code (Código de Comprimento de Dados)"),
    ("DTC", "Diagnostic Trouble Code (Código de Falha de Diagnóstico)"),
    ("ECU", "Electronic Control Unit (Unidade Eletrônica de Controle)"),
    ("FC", "Flow Control (controle de fluxo, no ISO-TP)"),
    ("FF", "First Frame (primeiro quadro, no ISO-TP)"),
    ("HIL", "Hardware-in-the-Loop (hardware no laço)"),
    ("HTTP", "Hypertext Transfer Protocol"),
    ("IoT", "Internet of Things (Internet das Coisas)"),
    ("ISO", "International Organization for Standardization"),
    ("ISO-TP", "Protocolo de transporte para CAN (ISO 15765-2)"),
    ("JSON", "JavaScript Object Notation"),
    ("MAP", "Manifold Absolute Pressure (Pressão Absoluta do Coletor)"),
    ("MIL", "Malfunction Indicator Lamp (Lâmpada Indicadora de Mau Funcionamento)"),
    ("MQTT", "Message Queuing Telemetry Transport"),
    ("OBD-II", "On-Board Diagnostics, segunda geração"),
    ("P2", "Tempo máximo de resposta da ECU a uma requisição (50 ms)"),
    ("PCI", "Protocol Control Information (informação de controle de protocolo)"),
    ("PID", "Parameter Identifier (Identificador de Parâmetro)"),
    ("QoS", "Quality of Service (Qualidade de Serviço)"),
    ("RF", "Requisito funcional"),
    ("RNF", "Requisito não funcional"),
    ("RPM", "Rotações por minuto"),
    ("RTT", "Round-Trip Time (Tempo de Ida e Volta)"),
    ("SAE", "Society of Automotive Engineers"),
    ("SF", "Single Frame (quadro único, no ISO-TP)"),
    ("SIL", "Software-in-the-Loop (software no laço)"),
    ("SQL", "Structured Query Language"),
    ("STFT", "Short-Term Fuel Trim (Ajuste de Combustível de Curto Prazo)"),
    ("STmin", "Separation Time minimum (tempo mínimo de separação entre quadros)"),
    ("TCP", "Transmission Control Protocol"),
    ("TWAI", "Two-Wire Automotive Interface"),
    ("VIN", "Vehicle Identification Number (Número de Identificação do Veículo)"),
]


def sh(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, check=True, capture_output=True, text=True, **kw)


# ---------------------------------------------------------------- pré-processamento
ESTADO = {"apendice": False}


def numerar_titulos(texto: str, contador: list[int]) -> tuple[str, list[tuple[int, str, str]]]:
    """Insere numeração nos títulos de capítulos e devolve (texto, entradas do sumário)."""
    entradas: list[tuple[int, str, str]] = []
    saida: list[str] = []
    em_codigo = False
    for linha in texto.split("\n"):
        if linha.startswith("```"):
            em_codigo = not em_codigo
        m = re.match(r"^(#{1,3}) (.+)$", linha) if not em_codigo else None
        if not m:
            saida.append(linha)
            continue
        nivel = len(m.group(1))
        titulo = m.group(2).strip()
        if nivel > 1 and ESTADO["apendice"]:
            entradas.append((2, "", titulo))
            saida.append(f"{'#' * nivel} {titulo} {{.unnumbered}}")
            continue
        if titulo.startswith("APÊNDICE") or titulo == "REFERÊNCIAS":
            ESTADO["apendice"] = titulo.startswith("APÊNDICE")
            entradas.append((1, "", titulo))
            saida.append(f"# {titulo} {{.unnumbered}}")
            continue
        contador[nivel - 1] += 1
        for i in range(nivel, 3):
            contador[i] = 0
        num = ".".join(str(c) for c in contador[:nivel])
        entradas.append((nivel, num, titulo))
        saida.append(f"{'#' * nivel} {num} {titulo}")
    return "\n".join(saida), entradas


def figuras_legenda_acima(texto: str) -> str:
    """ABNT: título da figura acima da imagem. Move a legenda para antes da imagem."""
    pad = re.compile(
        r"(!\[\]\([^)]*\)(?:\{[^}]*\})?)\n\n"
        r"(::: \{custom-style=\"LegendaFigura\"\}\n.*?\n:::)",
        re.S,
    )
    return pad.sub(lambda m: m.group(2) + "\n\n" + m.group(1), texto)


def legendas(texto: str, estilo: str) -> list[str]:
    pad = re.compile(
        r"::: \{custom-style=\"" + estilo + r"\"\}\n(.*?)\n:::", re.S
    )
    return [m.group(1).strip() for m in pad.finditer(texto)]


def p_toc(estilo: str, texto: str, pagina: int | None, recuo: int = 0) -> str:
    pg = "" if pagina is None else str(pagina)
    return (
        "```{=openxml}\n"
        f'<w:p><w:pPr><w:pStyle w:val="{estilo}"/><w:tabs><w:tab w:val="right" '
        f'w:leader="dot" w:pos="{LARGURA_TWIPS}"/></w:tabs>'
        f'<w:ind w:left="{recuo}" w:right="567" w:hanging="0"/></w:pPr>'
        f"<w:r><w:t xml:space=\"preserve\">{escape(texto)}</w:t></w:r><w:r><w:tab/></w:r>"
        f"<w:r><w:t>{pg}</w:t></w:r></w:p>\n```\n"
    )


def montar_markdown(paginas: dict[str, int]) -> str:
    contador = [0, 0, 0]
    ESTADO["apendice"] = False
    corpo_partes: list[str] = []
    entradas_todas: list[tuple[int, str, str]] = []
    for nome in CAPITULOS + POS:
        t = (AQUI / nome).read_text(encoding="utf-8")
        t = figuras_legenda_acima(t)
        t, ent = numerar_titulos(t, contador)
        entradas_todas += ent
        corpo_partes.append(t)
    corpo = "\n\n".join(corpo_partes)

    figs = legendas(corpo, "LegendaFigura")
    tabs = legendas(corpo, "LegendaTabela")

    def lista(itens: list[str]) -> str:
        return "".join(p_toc("ItemLista", i, paginas.get(i)) for i in itens)

    sumario = ""
    for nivel, num, titulo in entradas_todas:
        rot = f"{num} {titulo}".strip()
        sumario += p_toc(f"Sumario{nivel}", rot, paginas.get(rot), recuo=0)

    siglas = "| Sigla | Significado |\n| --- | --- |\n" + "\n".join(
        f"| {a} | {b} |" for a, b in SIGLAS
    )

    pre = (AQUI / "00-pre-textuais.md").read_text(encoding="utf-8")
    pre = pre.replace("<!--LISTA_FIGURAS-->", lista(figs))
    pre = pre.replace("<!--LISTA_TABELAS-->", lista(tabs))
    pre = pre.replace("<!--SIGLAS-->", siglas)
    pre = pre.replace("<!--SUMARIO-->", sumario)

    quebra_secao = (
        "\n```{=openxml}\n"
        '<w:p><w:pPr><w:sectPr><w:type w:val="nextPage"/>'
        '<w:pgSz w:w="11906" w:h="16838"/>'
        '<w:pgMar w:top="1701" w:right="1134" w:bottom="1134" w:left="1701" '
        'w:header="709" w:footer="709" w:gutter="0"/></w:sectPr></w:pPr></w:p>\n```\n\n'
    )
    md = pre + quebra_secao + corpo
    (AQUI / "build").mkdir(exist_ok=True)
    (AQUI / "build" / "tcc_completo.md").write_text(md, encoding="utf-8")
    return md, entradas_todas, figs, tabs  # type: ignore[return-value]


# ---------------------------------------------------------------- modelo de estilos
def _set_font(style, nome=FONTE, tam=12, negrito=None, italico=None, cor=None):
    f = style.font
    f.name = nome
    f.size = Pt(tam)
    if negrito is not None:
        f.bold = negrito
    if italico is not None:
        f.italic = italico
    f.color.rgb = cor or RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), nome)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def _par(style, alinhamento=None, esp=1.5, antes=0, depois=0, recuo=None, keep_next=None):
    pf = style.paragraph_format
    if alinhamento is not None:
        pf.alignment = alinhamento
    pf.line_spacing = esp
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    if recuo is not None:
        pf.first_line_indent = Cm(recuo)
    if keep_next is not None:
        pf.keep_with_next = keep_next


def criar_modelo(caminho: Path) -> None:
    ref = AQUI / "build" / "ref_pandoc.docx"
    dados = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"], check=True, capture_output=True
    ).stdout
    ref.write_bytes(dados)
    d = Document(str(ref))
    st = d.styles
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    L = WD_ALIGN_PARAGRAPH.LEFT
    C = WD_ALIGN_PARAGRAPH.CENTER

    def obter(nome, tipo=1):
        try:
            return st[nome]
        except KeyError:
            return st.add_style(nome, tipo)

    _set_font(st["Normal"])
    _par(st["Normal"], J, 1.5)
    for nome in ("Body Text", "First Paragraph"):
        s = obter(nome)
        _set_font(s)
        _par(s, J, 1.5, recuo=1.25)
    s = obter("Compact")
    _set_font(s)
    _par(s, J, 1.5)
    s = obter("Block Text")
    _set_font(s, tam=10)
    _par(s, J, 1.0)

    for i, (tam, caps) in enumerate([(12, True), (12, True), (12, False)], start=1):
        h = obter(f"Heading {i}")
        _set_font(h, tam=tam, negrito=True)
        _par(h, L, 1.5, antes=18 if i == 1 else 12, depois=12, keep_next=True)
        h.paragraph_format.first_line_indent = Cm(0)
        if i == 1:
            h.paragraph_format.page_break_before = True

    s = obter("Source Code", 1)
    _set_font(s, "Courier New", 9)
    _par(s, L, 1.0, antes=6, depois=6)
    s.paragraph_format.left_indent = Cm(0.5)
    try:
        v = st["Verbatim Char"]
        _set_font(v, "Courier New", 9)
    except KeyError:
        pass

    def novo(nome, tam=12, negrito=False, alin=C, esp=1.0, antes=0, depois=0, recuo=0.0,
             keep=None, italico=None):
        s = obter(nome)
        s.base_style = st["Normal"]
        _set_font(s, tam=tam, negrito=negrito, italico=italico)
        _par(s, alin, esp, antes, depois, recuo, keep)
        s.paragraph_format.left_indent = Cm(0)
        return s

    novo("Capa", negrito=True, antes=0, depois=14)
    novo("CapaTitulo", negrito=True, antes=90, depois=90)
    novo("CapaRodape", antes=110, depois=0)
    s = novo("Natureza", alin=J, antes=12, depois=8)
    s.paragraph_format.left_indent = Cm(8)
    novo("Assinatura", antes=10, depois=0)
    s = novo("Dedicatoria", alin=J, esp=1.5, antes=300)
    s.paragraph_format.left_indent = Cm(8)
    novo("TituloPre", negrito=True, esp=1.5, antes=0, depois=18)
    novo("Resumo", alin=J, esp=1.0, antes=0, depois=12)
    novo("LegendaFigura", tam=10, negrito=True, alin=L, antes=12, depois=4, keep=True)
    novo("LegendaTabela", tam=10, negrito=True, alin=L, antes=12, depois=4, keep=True)
    novo("Fonte", tam=10, alin=L, antes=4, depois=12)
    for n in (1, 2, 3):
        s = novo(f"Sumario{n}", alin=L, esp=1.15, antes=3 if n > 1 else 8, depois=0,
                 negrito=(n == 1))
        s.paragraph_format.left_indent = Cm(0.6 * (n - 1))
    novo("ItemLista", alin=L, esp=1.15, antes=4, depois=0)
    d.save(str(caminho))


# ---------------------------------------------------------------- pós-processamento
def _bordas(tbl, siglas: bool) -> None:
    tblPr = tbl._tbl.tblPr
    antigo = tblPr.find(qn("w:tblBorders"))
    if antigo is not None:
        tblPr.remove(antigo)
    b = OxmlElement("w:tblBorders")
    nomes = ["top", "bottom", "insideH"] if not siglas else []
    for n in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{n}")
        if n in nomes or (n in ("left", "right", "insideV") and not siglas and False):
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "808080")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def _larguras(tbl, ncol: int, textos: list[list[str]]) -> list[int]:
    pesos = []
    for c in range(ncol):
        mx = max(len(l[c]) if c < len(l) else 0 for l in textos)
        med = sum(min(len(l[c]), 60) if c < len(l) else 0 for l in textos) / max(len(textos), 1)
        palavra = max(
            (len(w) for l in textos if c < len(l) for w in l[c].split()), default=4
        )
        pesos.append(max(6.0, palavra * 1.15 + 1, 0.55 * min(mx, 40) + 0.45 * med))
    tot = sum(pesos)
    return [int(LARGURA_TWIPS * p / tot) for p in pesos]


def posprocessar(caminho: Path) -> None:
    d = Document(str(caminho))
    # seções: margens ABNT e cabeçalho com número só na parte textual
    for s in d.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(3), Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(2)
        s.header_distance = Cm(1.25)
    ultimo = d.sections[-1]
    ultimo.header.is_linked_to_previous = False
    hp = ultimo.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    run = hp.add_run()
    run.font.name = FONTE
    run.font.size = Pt(10)
    for tipo, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
    # a primeira seção (pré-textual) não tem cabeçalho
    d.sections[0].header.is_linked_to_previous = False

    for p in d.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)

    em_ref = False
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"):
            em_ref = p.text.strip() == "REFERÊNCIAS"
            continue
        if em_ref and p.text.strip():
            pf = p.paragraph_format
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = Cm(0)
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(12)

    for tbl in d.tables:
        textos = [[c.text for c in row.cells] for row in tbl.rows]
        siglas = bool(textos) and textos[0][:1] == ["Sigla"]
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        ncol = len(tbl.columns)
        if siglas:
            larg = [int(LARGURA_TWIPS * 0.18), int(LARGURA_TWIPS * 0.82)]
        else:
            larg = _larguras(tbl, ncol, textos)
        tblPr = tbl._tbl.tblPr
        for e in tblPr.findall(qn("w:tblW")):
            tblPr.remove(e)
        w = OxmlElement("w:tblW")
        w.set(qn("w:w"), str(sum(larg)))
        w.set(qn("w:type"), "dxa")
        tblPr.append(w)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        grid = tbl._tbl.tblGrid
        for gc, wv in zip(grid.findall(qn("w:gridCol")), larg):
            gc.set(qn("w:w"), str(wv))
        _bordas(tbl, siglas)
        for ri, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            if ri == 0 and not siglas:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for ci, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                for e in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(e)
                tw = OxmlElement("w:tcW")
                tw.set(qn("w:w"), str(larg[ci]))
                tw.set(qn("w:type"), "dxa")
                tcPr.insert(0, tw)
                if ri == 0 and not siglas:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "E7E6E6")
                    tcPr.append(shd)
                mar = OxmlElement("w:tcMar")
                for lado, v in (("top", 40), ("bottom", 40), ("left", 70), ("right", 70)):
                    m = OxmlElement(f"w:{lado}")
                    m.set(qn("w:w"), str(v))
                    m.set(qn("w:type"), "dxa")
                    mar.append(m)
                tcPr.append(mar)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    pf = p.paragraph_format
                    pf.first_line_indent = Cm(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(9 if not siglas else 11)
                        r.font.name = FONTE
                        if ri == 0 and not siglas:
                            r.font.bold = True
        # a legenda (parágrafo anterior) fica junto da tabela
    d.save(str(caminho))


# ---------------------------------------------------------------- paginação
def ler_paginas(pdf: Path, entradas, figs, tabs) -> dict[str, int]:
    txt = sh(["pdftotext", "-layout", str(pdf), "-"]).stdout
    paginas_txt = txt.split("\f")

    def norm(s: str) -> str:
        return re.sub(r"\s+", " ", s).strip()

    # primeira página da parte textual: título "1 INTRODUÇÃO" sem pontilhado
    inicio = 0
    for i, pg in enumerate(paginas_txt):
        for ln in pg.split("\n"):
            if norm(ln).startswith("1 INTRODUÇÃO") and "...." not in ln:
                inicio = i
                break
        if inicio:
            break
    res: dict[str, int] = {}

    def achar(rotulo: str, a_partir: int) -> int | None:
        alvo = norm(rotulo)[:38]
        for i in range(a_partir, len(paginas_txt)):
            for ln in paginas_txt[i].split("\n"):
                n = norm(ln)
                if n.startswith(alvo) and "...." not in n:
                    return i + 1
        return None

    for nivel, num, titulo in entradas:
        rot = f"{num} {titulo}".strip()
        pg = achar(rot, inicio)
        if pg:
            res[rot] = pg
    for lg in figs + tabs:
        pg = achar(lg, inicio)
        if pg:
            res[lg] = pg
    return res


def converter_pdf(docx: Path, destino: Path) -> Path:
    destino.mkdir(parents=True, exist_ok=True)
    sh(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(destino), str(docx)])
    return destino / (docx.stem + ".pdf")


def rodar(out: Path) -> None:
    (AQUI / "build").mkdir(exist_ok=True)
    modelo = AQUI / "build" / "modelo.docx"
    criar_modelo(modelo)
    paginas: dict[str, int] = {}
    nome = "TCC_Plataforma_Telemetria_OBD-II.docx"
    for passo in (1, 2, 3):
        md, entradas, figs, tabs = montar_markdown(paginas)  # type: ignore[misc]
        docx = AQUI / "build" / nome
        sh([
            "pandoc", str(AQUI / "build" / "tcc_completo.md"), "-f",
            "markdown+fenced_divs+raw_attribute+tex_math_dollars", "-o", str(docx),
            "--reference-doc", str(modelo), "--resource-path", str(AQUI),
            "--wrap=none",
        ])
        posprocessar(docx)
        pdf = converter_pdf(docx, AQUI / "build")
        novas = ler_paginas(pdf, entradas, figs, tabs)
        print(f"passo {passo}: {len(novas)} entradas paginadas", file=sys.stderr)
        if novas == paginas:
            break
        paginas = novas
    out.mkdir(parents=True, exist_ok=True)
    for f in (docx, pdf):
        (out / f.name).write_bytes(f.read_bytes())
    print("gerado:", out / docx.name, out / pdf.name)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=AQUI / "build")
    rodar(ap.parse_args().out)
